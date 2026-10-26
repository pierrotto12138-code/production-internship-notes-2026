# -*- coding: utf-8 -*-
"""在用户修改版报告里补入摘要正文与关键词，并按当前排版刷新目录页码缓存。
只改这两处，正文与其余格式保持用户版本不动。保存为原地覆盖（备份已另行创建）。"""
import copy
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
import content as C

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

PATH = str(ROOT / "报告" / "专业生产实习报告_已脱敏.docx")

# 目录条目 -> 新页码（按用户修改后的排版实测：物理页脚码）
NEW_PAGES = {
    "第一章实习概况": 1, "1.1实习基本情况": 1, "1.2实习目的": 1, "1.3实习任务与安排": 1,
    "第二章工程视觉": 3, "2.1实训平台与总体思路": 3, "2.2定位与标定": 3, "2.3数据格式化与脚本": 4,
    "2.4精密测量与通信": 5, "2.5缺陷检测与识别分类": 7, "2.6典型问题与工程经验": 9,
    "第三章机器人编程": 11, "3.1实训平台与总体安排": 11, "3.2工作站搭建与坐标系": 11,
    "3.3运动指令与IO控制": 13, "3.4RAPID程序设计与练习": 14, "3.5调试排错经验": 18,
    "第四章实习收获": 20, "4.1知识与技能收获": 20, "4.2工程方法与职业素养": 20, "4.3不足与今后方向": 20,
    "第五章致谢": 22,
}

doc = Document(PATH)
paras = list(doc.paragraphs)

def find_idx(pred):
    for i, p in enumerate(paras):
        if pred(p):
            return i
    raise ValueError("not found")

i_abs = find_idx(lambda p: p.text.strip() == "摘要")
i_tpl = find_idx(lambda p: p.text.strip().startswith("本次专业生产实习于2026年9月7日"))
tpl_el = paras[i_tpl]._p
print("摘要标题段 idx:", i_abs, "| 正文样式模板段 idx:", i_tpl)

def strip_bm(el):
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in el.findall(".//" + qn(tag)):
            bm.getparent().remove(bm)

def make_par(texts, drop_ind=False):
    """按正文段落模板生成新段；texts 为 (文本, 是否加粗) 列表"""
    el = copy.deepcopy(tpl_el)
    strip_bm(el)
    for r in el.findall(qn("w:r")):
        el.remove(r)
    if drop_ind:
        ppr = el.find(qn("w:pPr"))
        if ppr is not None:
            ind = ppr.find(qn("w:ind"))
            if ind is not None:
                ppr.remove(ind)
    for text, bold in texts:
        r = etree.SubElement(el, qn("w:r"))
        if bold:
            rpr = etree.SubElement(r, qn("w:rPr"))
            etree.SubElement(rpr, qn("w:b"))
        t = etree.SubElement(r, qn("w:t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
    return el

abs_text = C.BLOCKS[0][1]
kw_text = C.BLOCKS[1][1]
abs_el = make_par([(abs_text, False)])
kw_el = make_par([("关键词：", True), (kw_text, False)], drop_ind=True)

p_abs = paras[i_abs]._p
p_abs.addnext(kw_el)
p_abs.addnext(abs_el)
print("已插入摘要正文（%d 字）与关键词行" % len(abs_text))

# 刷新目录页码缓存
changed = 0
for p in doc.paragraphs:
    st = (p.style.name if p.style else "").lower()
    if not st.startswith("toc") or p.text.strip() == "目录":
        continue
    key = re.sub(r"\s+", "", p.text.rsplit("\t", 1)[0] if "\t" in p.text else p.text)
    if key not in NEW_PAGES:
        print("  ?? 无映射:", repr(p.text))
        continue
    new = NEW_PAGES[key]
    target = None
    for r in p._p.findall(".//" + qn("w:r")):
        if r.find(qn("w:tab")) is not None and r.find(qn("w:t")) is not None:
            target = r
    if target is None:
        print("  ?? 未找到页码 run:", repr(p.text))
        continue
    t = target.find(qn("w:t"))
    old = t.text
    if old != str(new):
        t.text = str(new)
        changed += 1
        print("  目录: %s  %s -> %s" % (key, old, new))
print("目录页码更新条数:", changed)

doc.save(PATH)
print("已保存:", PATH)
