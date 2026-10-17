# -*- coding: utf-8 -*-
"""把封面单独切成 1 页 1 节，使摘要/目录所在节只有 2 页，页脚显示 x/2。
用法: python split_front.py <docx> <正文总页数>
"""
import copy
import sys

from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
body_pages = int(sys.argv[2])

doc = Document(path)
ps = doc.paragraphs
i_abs = next(i for i, p in enumerate(ps) if p.text.strip() == "摘要")
i_h1 = next(i for i, p in enumerate(ps) if p.text.strip() == "实习概况")
prev = ps[i_abs - 1]          # 摘要前一段（原含分页符）
sec_holder = ps[i_h1 - 1]     # 承载“封面+摘要+目录”这一节 sectPr 的空段

pPr = sec_holder._p.find(qn("w:pPr"))
sectPr = pPr.find(qn("w:sectPr"))
assert sectPr is not None, "未找到前置部分 sectPr"

# 1) 封面节：把原 sectPr 复制一份，放进 prev 段，作为封面节的节属性
cover_sect = copy.deepcopy(sectPr)
# 删除分页符（改由分节符换页，避免空白页）；分页符可能位于 prev 段或摘要段
for _par in (prev._p, ps[i_abs]._p):
    for br in list(_par.iter(qn("w:br"))):
        if br.get(qn("w:type")) == "page":
            br.getparent().remove(br)
prev_pPr = prev._p.get_or_add_pPr()
prev_pPr.append(cover_sect)

# 2) 摘要+目录节（即原 sectPr）：去掉 titlePg（否则首页页脚为空，摘要不显示页码）
tp = sectPr.find(qn("w:titlePg"))
if tp is not None:
    sectPr.remove(tp)
# 页码从 1 重新开始
pgn = sectPr.find(qn("w:pgNumType"))
if pgn is None:
    from docx.oxml import OxmlElement
    pgn = OxmlElement("w:pgNumType")
    sectPr.insert(0, pgn)
pgn.set(qn("w:start"), "1")

doc.save(path)

# 3) 修正各节页脚“本节页数”域的缓存值（LibreOffice 不重算，需写正确显示值）
doc = Document(path)
secs = doc.sections
print("分节数:", len(secs))
want = [2, body_pages]

def set_sectpages_cache(part, val):
    if part is None:
        return 0
    n = 0
    for p in part.paragraphs:
        runs = p._p.findall(qn("w:r"))
        for idx, r in enumerate(runs):
            instr = r.find(qn("w:instrText"))
            if instr is not None and instr.text and "SECTIONPAGES" in instr.text:
                for r2 in runs[idx + 1:]:
                    t = r2.find(qn("w:t"))
                    if t is not None:
                        t.text = str(val)
                        n += 1
                        break
    return n

# secs[0]=封面(不显示), secs[1]=摘要+目录, secs[2]=正文
for idx, sec in enumerate(secs):
    val = 1 if idx == 0 else want[idx - 1]
    c = set_sectpages_cache(sec.footer, val)
    c += set_sectpages_cache(sec.first_page_footer, val)
    print(f"  节{idx+1}: 更新字段缓存 {c} 处 -> {val}")
doc.save(path)
print("saved:", path)
