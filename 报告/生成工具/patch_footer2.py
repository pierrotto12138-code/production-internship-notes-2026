# -*- coding: utf-8 -*-
"""把各节页脚里的“总页数”域 NUMPAGES 改为“本节页数”域 SECTIONPAGES，
并把该字段的缓存值改为正确数字，使正文页脚显示 当前页/正文总页数。
用法: python patch_footer2.py <docx> <节1页数> <节2页数>
"""
import sys

from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
dens = [int(sys.argv[2]), int(sys.argv[3])]

doc = Document(path)
assert len(doc.sections) == 2, len(doc.sections)
total = 0
for i, sec in enumerate(doc.sections):
    den = dens[i]
    for part in (sec.footer, sec.first_page_footer):
        try:
            part.paragraphs
        except Exception:
            continue
        for p in part.paragraphs:
            runs = p._p.findall(qn("w:r"))
            for idx, r in enumerate(runs):
                instr = r.find(qn("w:instrText"))
                if instr is not None and instr.text and "NUMPAGES" in instr.text:
                    instr.text = instr.text.replace("NUMPAGES", "SECTIONPAGES")
                    total += 1
                    # 紧随其后的第一个含 w:t 的 run 即该字段的缓存值
                    for r2 in runs[idx + 1:]:
                        t = r2.find(qn("w:t"))
                        if t is not None:
                            t.text = str(den)
                            break
print("替换字段数:", total)
doc.save(path)
print("saved:", path)
