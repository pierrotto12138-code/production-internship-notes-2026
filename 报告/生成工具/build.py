# -*- coding: utf-8 -*-
"""把 content.py 的内容装配进实习报告模板（python-docx 版，脚本自身不写文件：
读取/保存由 python-docx 完成；手抄稿 md 打印到 stdout 由调用方重定向）。
- 封面填空（班级/姓名/学号）
- 第一章改名“实习概况”，新增“实习收获”章（克隆模板 H1，保留“第X章”自动编号）
- 摘要/正文/插图/程序清单按模板结构插入
- 目录缓存条目重建（保留 TOC 域，Word 中可 F9 刷新）
页码两遍法：渲染后把页码表写入 content.py 的 PAGES，再跑一遍。
"""
import copy
import importlib
import struct
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))

_mod = sys.argv[1] if len(sys.argv) > 1 else "content"
C = importlib.import_module(_mod)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

TPL = str(HERE / "实习报告模板 (2)(1).docx")
OUT_DOCX = sys.argv[2] if len(sys.argv) > 2 else str(ROOT / "报告" / "专业生产实习报告_×××_2305××××××.docx")
IMG_BASE = str(ROOT) + "/"

pages = getattr(C, "PAGES", {})

doc = Document(TPL)
paras = list(doc.paragraphs)
# 模板：49 段（目录 6 段被包在 w:sdt 内不计入）。
# dp30=引言(H1) dp31=实习目的 dp32=1.2实习任务 dp33=空 dp34=工程视觉(H1)
# dp45=机器人编程(H1) dp46/47=空 dp48=致谢(H1)
assert len(paras) == 49, len(paras)
for _i, _t in ((30, "引言"), (32, "1.2实习任务"), (34, "工程视觉"), (45, "机器人编程"), (48, "致谢")):
    assert paras[_i].text.strip().startswith(_t), (_i, paras[_i].text)

# 提前留取样式模板（后面会删除这些段落）
H1_TPL = copy.deepcopy(paras[30]._p)   # 一级标题：保留 numId=4（第%1章 自动编号）
H2_TPL = copy.deepcopy(paras[32]._p)   # 二级标题：“1.2实习任务” 段（手工编号样式）

# ---------- 工具 ----------
def add_after(par):
    new_p = par._p.makeelement(qn("w:p"), {})
    par._p.addnext(new_p)
    return Paragraph(new_p, par._parent)

def strip_bookmarks(el):
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in el.findall(".//" + qn(tag)):
            bm.getparent().remove(bm)

def set_par_text(par, text):
    runs = list(par.runs)
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def clone_h_par(tpl_el, text, anchor_par, before=False):
    el = copy.deepcopy(tpl_el)
    strip_bookmarks(el)
    for r in el.findall(qn("w:r")):
        el.remove(r)
    p = Paragraph(el, anchor_par._parent)
    p.add_run(text)
    if before:
        anchor_par._p.addprevious(el)
    else:
        anchor_par._p.addnext(el)
    return p

def h2_par(text, after_par):
    return clone_h_par(H2_TPL, text, after_par)

def shade_par(p, fill="F2F2F2"):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    # 依 OOXML 顺序：shd 应位于 tabs/spacing/ind/jc/rPr 之前
    for tag in ("w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr"):
        el = ppr.find(qn(tag))
        if el is not None:
            el.addprevious(shd)
            return
    ppr.append(shd)

def body_par(text, after_par):
    p = add_after(after_par)
    p.style = doc.styles["Normal"]
    p.paragraph_format.first_line_indent = Pt(24)   # 首行缩进两字符（12pt 字号）
    p.add_run(text)
    return p

def kw_par(text, after_par):
    p = add_after(after_par)
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_before = Pt(6)
    r1 = p.add_run("关键词：")
    r1.bold = True
    p.add_run(text)
    return p

def img_par(src, w_cm, after_par):
    p = add_after(after_par)
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(src, width=Cm(w_cm))
    return p

def caption_par(text, after_par, keep_next=False):
    p = add_after(after_par)
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p

def code_par(lines, after_par):
    p = add_after(after_par)
    p.style = doc.styles["Normal"]
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    shade_par(p)
    for i, ln in enumerate(lines):
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        if i < len(lines) - 1:
            r.add_break()
    return p

def img_size(p):
    with open(p, "rb") as f:
        d = f.read(64)
    if d[:8] == b"\x89PNG\r\n\x1a\n":
        w, h = struct.unpack(">II", d[16:24])
        return w, h
    with open(p, "rb") as f:
        data = f.read()
    i = 2
    while i < len(data) - 9:
        if data[i] == 0xFF and data[i + 1] in (0xC0, 0xC1, 0xC2, 0xC3):
            h, w = struct.unpack(">HH", data[i + 5:i + 9])
            return w, h
        elif data[i] == 0xFF and data[i + 1] not in (0xD8, 0x01) and not (0xD0 <= data[i + 1] <= 0xD9):
            i += 2 + struct.unpack(">H", data[i + 2:i + 4])[0]
        else:
            i += 1
    raise ValueError(p)

def display_cm(px_w, px_h):
    r = px_h / px_w
    if r > 1.15:
        w = 8.0
    elif r > 0.85:
        w = 10.5
    else:
        w = 13.0
    w = min(w, px_w / 96 * 2.54)   # 不放大
    return w, w * r

# ---------- 封面填空 ----------
def fill_cover(par, value):
    units = sum(2 if ord(c) > 0x2E7F else 1 for c in value)
    pad = max(4, 27 - units)
    left = pad // 2
    txt = " " * left + value + " " * (pad - left)
    u_runs = [r for r in par.runs if r.font.underline]
    u_runs[0].text = txt
    for r in u_runs[1:]:
        r._r.getparent().remove(r._r)

fill_cover(paras[8], dict(C.META)["专业班级"])
fill_cover(paras[9], dict(C.META)["姓名"])
fill_cover(paras[10], dict(C.META)["学号"])

# ---------- 目录（缓存段在 w:sdt 内） ----------
from lxml import etree

def toc_run_text(text):
    r = etree.Element(qn("w:r"))
    rpr = etree.SubElement(r, qn("w:rPr"))
    st = etree.SubElement(rpr, qn("w:rStyle"))
    st.set(qn("w:val"), "Style14")
    t = etree.SubElement(r, qn("w:t"))
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return r

def toc_run_tabpage(page):
    r = etree.Element(qn("w:r"))
    rpr = etree.SubElement(r, qn("w:rPr"))
    st = etree.SubElement(rpr, qn("w:rStyle"))
    st.set(qn("w:val"), "Style14")
    etree.SubElement(r, qn("w:tab"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = str(page)
    return r

def rebuild_toc_runs(par, text, page):
    """按模板结构重建 hyperlink 内 run：TOC2=编号+tab+标题+tab+页码；TOC1=标题+tab+页码"""
    hl = par._p.find(qn("w:hyperlink"))
    for r in hl.findall(qn("w:r")):
        hl.remove(r)
    parts = text.split(" ", 1)
    if par._p.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "TOC2" and len(parts) == 2:
        hl.append(toc_run_text(parts[0]))
        mid = etree.Element(qn("w:r"))
        rpr = etree.SubElement(mid, qn("w:rPr"))
        st = etree.SubElement(rpr, qn("w:rStyle"))
        st.set(qn("w:val"), "Style14")
        etree.SubElement(mid, qn("w:tab"))
        hl.append(mid)
        hl.append(toc_run_text(parts[1]))
    else:
        hl.append(toc_run_text(text))
    hl.append(toc_run_tabpage(page))

sdt_content = doc.element.body.find(qn("w:sdt") + "/" + qn("w:sdtContent"))
toc_pars = sdt_content.findall(qn("w:p"))
assert len(toc_pars) == 6, len(toc_pars)
entries = [(lvl, txt, pages.get(txt.replace(" ", ""), "1")) for lvl, txt in C.TOC_ENTRIES]
rebuild_toc_runs(Paragraph(toc_pars[0], None), entries[0][1], entries[0][2])
rebuild_toc_runs(Paragraph(toc_pars[-1], None), entries[-1][1], entries[-1][2])
anchor_el = toc_pars[0]
for lvl, txt, page in entries[1:-1]:
    tpl_el = toc_pars[3] if lvl == "TOC1" else toc_pars[1]
    el = copy.deepcopy(tpl_el)
    for bm_tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in el.findall(".//" + qn(bm_tag)):
            bm.getparent().remove(bm)
    rebuild_toc_runs(Paragraph(el, None), txt, page)
    anchor_el.addnext(el)
    anchor_el = el
# 删除原缓存条目 1~4（新序列已含全部条目）
for el_old in toc_pars[1:5]:
    sdt_content.remove(el_old)

# ---------- 正文分组 ----------
h1_idx = [i for i, (k, *_x) in enumerate(C.BLOCKS) if k == "h1"]
assert len(h1_idx) == 5, h1_idx
g_concept = C.BLOCKS[h1_idx[0] + 1:h1_idx[1]]
g_vis = C.BLOCKS[h1_idx[1] + 1:h1_idx[2]]
g_rob = C.BLOCKS[h1_idx[2] + 1:h1_idx[3]]
g_recv = C.BLOCKS[h1_idx[3] + 1:h1_idx[4]]
g_thx = C.BLOCKS[h1_idx[4] + 1:]

def emit_blocks(blocks, after_par):
    anchor = after_par
    last = after_par
    for kind, *rest in blocks:
        if kind == "h2":
            last = h2_par(rest[0], anchor)
        elif kind in ("p", "thanks"):
            last = body_par(rest[0], anchor)
        elif kind == "img":
            path, cap = rest
            src = IMG_BASE + path.replace("/", "\\")
            w_px, h_px = img_size(src)
            w_cm, _ = display_cm(w_px, h_px)
            pic = img_par(src, w_cm, anchor)
            last = caption_par(cap, pic)
        elif kind == "code":
            code_text, cap = rest
            cap_par = caption_par(cap, anchor, keep_next=True)
            last = code_par(code_text.split("\n"), cap_par)
        elif kind in ("abstract", "keywords"):
            pass
        else:
            raise ValueError(kind)
        anchor = last
    return last

# ---------- 组装（先改标题与删空段，再自后向前插入） ----------
# 第一章改名
set_par_text(paras[30], "实习概况")
# 删除模板遗留：dp31（实习目的）、dp32（1.2实习任务）、dp33（空段）
for i in (31, 32, 33):
    el = paras[i]._p
    el.getparent().remove(el)
# 删除机器人章与致谢之间的两个空段
for i in (46, 47):
    el = paras[i]._p
    el.getparent().remove(el)

# 第一章正文：插入在 dp30（实习概况）之后
emit_blocks(g_concept, paras[30])
# 第二、三章
emit_blocks(g_vis, paras[34])
emit_blocks(g_rob, paras[45])
# 第四章：新建 H1“实习收获”（插在致谢之前），正文接在其后
h1_recv = clone_h_par(H1_TPL, "实习收获", paras[48], before=True)
emit_blocks(g_recv, h1_recv)
# 致谢正文
emit_blocks(g_thx, paras[48])
# 摘要标题（dp27）后：摘要正文 + 关键词行
kw_par(C.BLOCKS[1][1], body_par(C.BLOCKS[0][1], paras[27]))

doc.save(OUT_DOCX)
print("docx ->", OUT_DOCX, file=sys.stderr)
print("正文汉字数(不含程序清单):", C.cjk_count(), file=sys.stderr)

# ---------- 手抄稿 md 打到 stdout（调用方重定向） ----------
CH = ["第一章", "第二章", "第三章", "第四章", "第五章"]
code_n = sum(1 for blk in C.BLOCKS if blk[0] == "code")
out_name = OUT_DOCX.rsplit("/", 1)[-1]
md = ["# 《专业生产实习》报告（手抄稿）", "",
      "> 依据本仓库学习记录与实习日志整理（2026-09-28），未验证内容不写入；正文约 %d 字（另有 %d 个程序清单）。" % (C.cjk_count(), code_n),
      "> 配套正式电子版：`%s`（图号、清单号一致，手抄时对照插图即可）。" % out_name,
      ""]
chap = -1
for kind, *rest in C.BLOCKS:
    if kind == "abstract":
        md += ["**摘要**", "", rest[0], ""]
    elif kind == "keywords":
        md += ["**关键词**：" + rest[0], ""]
    elif kind == "h1":
        chap += 1
        md += ["## " + CH[chap] + " " + rest[0], ""]
    elif kind == "h2":
        md += ["### " + rest[0], ""]
    elif kind in ("p", "thanks"):
        md += [rest[0], ""]
    elif kind == "img":
        path, cap = rest
        md += ["![" + cap + "](" + path + ")", "", "**" + cap + "**", ""]
    elif kind == "code":
        code_text, cap = rest
        md += ["**" + cap + "**", "", "```rapid", code_text, "```", ""]
print("\n".join(md))
